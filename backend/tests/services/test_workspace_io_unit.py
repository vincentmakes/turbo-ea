"""Unit tests for the workspace_io bundle plumbing — no database required."""

from __future__ import annotations

import openpyxl

from app.services.workspace_io import bundle as bundle_io
from app.services.workspace_io import schema
from app.services.workspace_io.secrets import strip_secrets

# ---------------------------------------------------------------------------
# Reference encoding (the round-trip risk)
# ---------------------------------------------------------------------------


def test_encode_decode_path_roundtrip_simple():
    segs = ["Europe", "Germany", "Berlin"]
    encoded = schema.encode_path(segs)
    assert schema.split_escaped_path(encoded) == segs


def test_encode_decode_path_with_slash_and_backslash_in_names():
    # Names containing the separator characters must survive intact.
    segs = ["SAP S/4HANA", "CI/CD Pipelines", "Back\\Office"]
    encoded = schema.encode_path(segs)
    # No raw " / " separators introduced by the names themselves.
    assert schema.split_escaped_path(encoded) == segs


def test_build_ref_string_matches_resolver_format():
    ref = schema.build_ref_string(["Org", "Unit"], "App/Name")
    # Parent segments + escaped name joined by the path separator.
    assert ref == "Org / Unit / App\\/Name"
    assert schema.split_escaped_path(ref) == ["Org", "Unit", "App/Name"]


def test_empty_path():
    assert schema.split_escaped_path("") == []
    assert schema.split_escaped_path("   ") == []
    assert schema.encode_path([]) == ""


# ---------------------------------------------------------------------------
# Secret stripping
# ---------------------------------------------------------------------------


def test_strip_secrets_removes_known_credentials():
    general = {
        "currency": "EUR",
        "sso": {"enabled": True, "client_id": "abc", "client_secret": "enc:SHHH"},
        "ai": {"enabled": True, "model": "gpt", "apiKey": "enc:KEY"},
    }
    email = {
        "smtp_host": "mail",
        "smtp_user": "u",
        "smtp_password": "enc:PW",
        "method": "graph_api",
        "oauth_client_id": "client-1",
        "oauth_client_secret": "enc:CS",
        "service_account_json": "enc:SA",
    }
    g, e = strip_secrets(general, email)
    assert g["currency"] == "EUR"
    assert "client_secret" not in g["sso"]
    assert g["sso"]["client_id"] == "abc"
    assert "apiKey" not in g["ai"]
    assert g["ai"]["model"] == "gpt"
    assert "smtp_password" not in e
    # New OAuth email secrets are stripped; non-secret fields survive.
    assert "oauth_client_secret" not in e
    assert "service_account_json" not in e
    assert e["smtp_host"] == "mail"
    assert e["method"] == "graph_api"
    assert e["oauth_client_id"] == "client-1"


def test_strip_secrets_scrubs_arbitrary_encrypted_value():
    general = {"nested": {"weird_secret": "enc:LEAK", "ok": "plain"}}
    g, _ = strip_secrets(general, {})
    assert "weird_secret" not in g["nested"]
    assert g["nested"]["ok"] == "plain"


def test_strip_secrets_does_not_mutate_input():
    general = {"sso": {"client_secret": "enc:x"}}
    strip_secrets(general, {})
    assert general["sso"]["client_secret"] == "enc:x"


# ---------------------------------------------------------------------------
# Cell (de)serialisation + zip pack/unpack
# ---------------------------------------------------------------------------


def test_cell_json_roundtrip():
    value = {"b": 2, "a": [1, 2, 3], "n": None}
    cell = bundle_io.to_cell(value, is_json=True)
    assert isinstance(cell, str)
    assert bundle_io.from_cell(cell, is_json=True) == value


def test_cell_scalar_passthrough():
    assert bundle_io.to_cell(True, is_json=False) is True
    assert bundle_io.to_cell(None, is_json=False) is None
    assert bundle_io.from_cell("hi", is_json=False) == "hi"


def test_pack_unpack_roundtrip():
    wb = openpyxl.Workbook(write_only=True)
    bundle_io.write_sheet(wb, "Demo", ["a", "b"], [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}])
    import io

    buf = io.BytesIO()
    wb.save(buf)
    raw = bundle_io.pack({"format_version": "1"}, buf.getvalue(), {"branding/logo.png": b"PNG"})
    manifest, workbook_bytes, assets = bundle_io.unpack(raw)
    assert manifest["format_version"] == "1"
    assert assets["branding/logo.png"] == b"PNG"
    sheets = bundle_io.read_workbook(workbook_bytes)
    assert sheets["Demo"] == [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]


def test_unpack_rejects_non_zip():
    import pytest

    from app.services.workspace_io.bundle import BundleFormatError

    with pytest.raises(BundleFormatError):
        bundle_io.unpack(b"not a zip")


def test_oversized_cell_survives_via_overflow_asset():
    """A cell beyond Excel's 32,767-char limit is offloaded to an asset and
    restored verbatim on parse (openpyxl would otherwise truncate it)."""
    import io
    import json

    big = json.dumps({"blob": "y" * 50000})
    assert len(big) > 32767

    wb = openpyxl.Workbook(write_only=True)
    assets: dict[str, bytes] = {}
    bundle_io.write_sheet(wb, "Big", ["v"], [{"v": big}], assets)
    assert any(k.startswith("overflow/") for k in assets)

    buf = io.BytesIO()
    wb.save(buf)
    raw = bundle_io.pack({"format_version": "1"}, buf.getvalue(), assets)

    parsed = bundle_io.parse_bundle(raw)
    restored = parsed.rows("Big")[0]["v"]
    assert restored == big  # full value, not truncated
    assert json.loads(restored)["blob"] == "y" * 50000


def test_merge_settings_never_writes_incoming_secrets():
    """A hand-edited/malicious bundle carrying a secret must not land it —
    neither overwriting the target's value nor creating one the target lacked
    (it would be stored unencrypted)."""
    from app.services.workspace_io.applier import _merge_settings
    from app.services.workspace_io.secrets import EMAIL_SECRET_PATHS

    target = {"smtp_host": "old-host", "smtp_password": "enc:KEEP"}
    incoming = {
        "smtp_host": "new-host",
        "method": "graph_api",
        "smtp_password": "plaintext-attack",
        "oauth_client_secret": "plaintext-attack",
        "service_account_json": '{"private_key": "attack"}',
    }
    merged = _merge_settings(target, incoming, EMAIL_SECRET_PATHS)
    assert merged["smtp_host"] == "new-host"
    assert merged["method"] == "graph_api"
    assert merged["smtp_password"] == "enc:KEEP"  # target's own value preserved
    assert "oauth_client_secret" not in merged  # never created from a bundle
    assert "service_account_json" not in merged


# ---------------------------------------------------------------------------
# Importing from a path: the assets tree stays on disk until it is asked for
# ---------------------------------------------------------------------------


def _demo_bundle_bytes(assets: dict[str, bytes] | None = None) -> bytes:
    import io

    wb = openpyxl.Workbook(write_only=True)
    bundle_io.write_sheet(wb, "Demo", ["a"], [{"a": 1}])
    buf = io.BytesIO()
    wb.save(buf)
    return bundle_io.pack(
        {"format_version": "1"}, buf.getvalue(), assets or {"branding/logo.png": b"PNG"}
    )


def test_parse_bundle_from_a_path_reads_assets_lazily(tmp_path):
    """An import parses from disk, so the applier can pull one asset at a time.

    A workspace's assets are its attachments, diagrams and branding — most of a
    large bundle. Inflating them all at parse time is what made a big import a
    memory problem.
    """
    path = tmp_path / "bundle.zip"
    path.write_bytes(_demo_bundle_bytes({"branding/logo.png": b"PNG", "files/a.bin": b"DATA"}))

    with bundle_io.parse_bundle(path) as parsed:
        assert isinstance(parsed.assets, bundle_io.ZipAssetStore)
        assert parsed.rows("Demo") == [{"a": 1}]
        assert "files/a.bin" in parsed.assets
        assert "files/missing.bin" not in parsed.assets
        assert parsed.assets.get("files/a.bin") == b"DATA"
        assert parsed.assets.get("files/missing.bin") is None
        assert parsed.assets.keys() == ["branding/logo.png", "files/a.bin"]

    # Past the context manager the zip is closed; asking for more is a bug in
    # the caller, and says so rather than returning a silent None.
    import pytest

    with pytest.raises(RuntimeError, match="closed"):
        parsed.assets.get("files/a.bin")


def test_parse_bundle_still_accepts_bytes(tmp_path):
    """Content packs and the tests build bundles in memory; that must keep working."""
    parsed = bundle_io.parse_bundle(_demo_bundle_bytes())
    assert parsed.rows("Demo") == [{"a": 1}]
    assert parsed.assets.get("branding/logo.png") == b"PNG"
    parsed.close()  # closes the in-memory zip; nothing on disk


def test_a_plain_dict_satisfies_the_asset_store_contract():
    # This is what lets build_content_bundle keep passing a dict.
    bundle = bundle_io.WorkspaceBundle(manifest={}, sheets={}, assets={"x": b"y"})
    assert bundle.assets.get("x") == b"y"
    assert "x" in bundle.assets
    assert list(bundle.assets.keys()) == ["x"]
    bundle.close()


def test_an_overflow_cell_is_restored_when_parsing_from_a_path(tmp_path):
    import io
    import json

    big = json.dumps({"blob": "z" * 50000})
    wb = openpyxl.Workbook(write_only=True)
    assets: dict[str, bytes] = {}
    bundle_io.write_sheet(wb, "Big", ["v"], [{"v": big}], assets)
    buf = io.BytesIO()
    wb.save(buf)
    path = tmp_path / "overflow.zip"
    path.write_bytes(bundle_io.pack({"format_version": "1"}, buf.getvalue(), assets))

    with bundle_io.parse_bundle(path) as parsed:
        assert parsed.rows("Big")[0]["v"] == big


def test_a_member_declaring_an_implausible_size_is_refused(tmp_path, monkeypatch):
    """The central directory is checked before a single member is read."""
    import pytest

    from app.services.workspace_io.bundle import BundleFormatError

    # Above the manifest and workbook, below the asset — so it is the asset
    # that trips, which is what the cap is for.
    monkeypatch.setattr(bundle_io, "MAX_ASSET_BYTES", 100_000)
    path = tmp_path / "big-member.zip"
    path.write_bytes(_demo_bundle_bytes({"files/big.bin": b"x" * 200_000}))

    with pytest.raises(BundleFormatError, match="files/big.bin"):
        bundle_io.parse_bundle(path)


def test_a_bundle_whose_members_sum_too_high_is_refused(tmp_path, monkeypatch):
    import pytest

    from app.services.workspace_io.bundle import BundleFormatError

    monkeypatch.setattr(bundle_io, "MAX_BUNDLE_UNCOMPRESSED_BYTES", 500_000)
    path = tmp_path / "bomb.zip"
    path.write_bytes(_demo_bundle_bytes({f"files/{i}.bin": b"x" * 100_000 for i in range(8)}))

    with pytest.raises(BundleFormatError, match="refused"):
        bundle_io.parse_bundle(path)


def test_a_corrupt_file_on_disk_is_reported_as_a_bad_bundle(tmp_path):
    import pytest

    from app.services.workspace_io.bundle import BundleFormatError

    path = tmp_path / "junk.zip"
    path.write_bytes(b"not a zip at all")
    with pytest.raises(BundleFormatError, match="not a valid .zip"):
        bundle_io.parse_bundle(path)
