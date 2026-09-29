"""Unit tests for the LeanIX xlsx exporter — the mirror of ``xlsx_parser.py``.

Everything here runs without a database: a :class:`MigrationSnapshot`
is built by hand in Turbo EA vocabulary, written with
:func:`build_workbook`, and read back either directly with openpyxl or
through the importer's own :func:`parse_xlsx`, so a Turbo EA → LeanIX
→ Turbo EA trip is pinned end to end.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.services.migration.snapshot import (
    Comment,
    Document,
    MetamodelField,
    MetamodelRelationType,
    MetamodelType,
    MigrationSnapshot,
    Relation,
    SourceEntity,
    Subscription,
    Tag,
)
from app.services.migration.sources.leanix.adapter import LeanixSource
from app.services.migration.sources.leanix.mappings import (
    EXPORT_FLIP_DIRECTION,
    EXPORT_RELATION_MAPPING,
    EXPORT_TYPE_MAPPING,
    FLIP_DIRECTION,
    RELATION_MAPPING,
    TYPE_MAPPING,
)
from app.services.migration.sources.leanix.xlsx_exporter import build_workbook
from app.services.migration.sources.leanix.xlsx_parser import is_xlsx_payload, parse_xlsx

# ---------------------------------------------------------------------------
# Fixture snapshot
# ---------------------------------------------------------------------------


def _snapshot() -> MigrationSnapshot:
    entities = [
        SourceEntity(
            source_id="app-1",
            type="Application",
            name="Salesforce",
            display_name="Salesforce",
            category="businessApplication",
            description="CRM",
            lifecycle={"plan": "2019-01-01", "active": "2020-01-01"},
            tags=["tag-emea", "tag-pilot"],
            custom_fields={
                "costTotalAnnual": 1200.5,
                "hostingType": "cloud",
                "regions": ["emea", "apac"],
                "externalId": "LX-42",
            },
            quality_seal="APPROVED",
            completion=0.85,
            status="ACTIVE",
            raw={"createdAt": "2024-01-01T00:00:00+00:00"},
        ),
        SourceEntity(
            source_id="app-2",
            type="Application",
            name="New CRM",
            display_name="New CRM",
            status="ACTIVE",
        ),
        SourceEntity(
            source_id="bc-1",
            type="BusinessCapability",
            name="Sales",
            display_name="Sales",
            status="ACTIVE",
        ),
        SourceEntity(
            source_id="bc-2",
            type="BusinessCapability",
            name="Lead Mgmt",
            display_name="Sales / Lead Mgmt",
            parent_id="bc-1",
            status="ACTIVE",
        ),
        SourceEntity(
            source_id="proc-1",
            type="BusinessProcess",
            name="Order to Cash",
            display_name="Order to Cash",
            status="ACTIVE",
        ),
        SourceEntity(
            source_id="org-1",
            type="Organization",
            name="Sales EMEA",
            display_name="Sales EMEA",
            category="team",
            status="ACTIVE",
        ),
        SourceEntity(
            source_id="srv-1",
            type="Server",  # admin-created type, no LeanIX equivalent
            name="db-01",
            display_name="db-01",
            status="ACTIVE",
        ),
    ]
    relations = [
        Relation("rel-1", "relAppToBC", "app-1", "bc-2", attributes={}),
        # Turbo EA: source succeeds target → app-2 succeeds app-1.
        Relation("rel-2", "relAppSuccessor", "app-2", "app-1", attributes={}),
        Relation("rel-3", "relProcessToApp", "proc-1", "app-1", attributes={}),
        Relation(
            "rel-4",
            "relOrgToApp",
            "org-1",
            "app-1",
            attributes={"usageType": "owner", "description": "Owning unit"},
        ),
        Relation("rel-5", "relServerToApp", "srv-1", "app-1", attributes={}),
    ]
    subscriptions = [
        Subscription("sub-1", "app-1", "owner@example.com", "Owner", "Responsible", "RESPONSIBLE"),
        Subscription("sub-2", "app-1", "a@example.com", "A", "Observer", "OBSERVER"),
        Subscription("sub-3", "app-1", "b@example.com", "B", "Observer", "OBSERVER"),
    ]
    tags = [
        Tag("tag-emea", "EMEA", "Region", "MULTIPLE", "#ff0000"),
        Tag("tag-apac", "APAC", "Region", "MULTIPLE", "#00ff00"),
        Tag(
            "tag-pilot",
            "Pilot",
            "Stage",
            "SINGLE",
            None,
            group_restrict_to_types=["Application"],
        ),
    ]
    documents = [Document("doc-1", "app-1", "Runbook", "https://wiki.example.com/sf")]
    comments = [
        Comment(
            "cmt-1",
            "app-1",
            "owner@example.com",
            "Renewal due in Q3",
            created_at=datetime(2024, 3, 1, 12, 0),
        )
    ]
    metamodel_types = [
        MetamodelType(
            "Application",
            False,
            fields=[
                MetamodelField("Application", "costTotalAnnual", "Total annual cost", "cost"),
                MetamodelField(
                    "Application",
                    "hostingType",
                    "Hosting type",
                    "single_select",
                    options=[
                        {"key": "cloud", "label": "Cloud"},
                        {"key": "onprem", "label": "On premise"},
                    ],
                ),
                MetamodelField(
                    "Application",
                    "regions",
                    "Regions",
                    "multiple_select",
                    options=[{"key": "emea", "label": "EMEA"}, {"key": "apac", "label": "APAC"}],
                ),
            ],
            subtypes=["businessApplication", "microservice"],
        ),
        MetamodelType("BusinessCapability", False),
        MetamodelType("BusinessProcess", False),
        MetamodelType("Organization", False, subtypes=["team"]),
        MetamodelType("Server", True),
    ]
    metamodel_relation_types = [
        MetamodelRelationType("relAppToBC", "Application", "BusinessCapability", "supports"),
        MetamodelRelationType("relAppSuccessor", "Application", "Application", "succeeds"),
        MetamodelRelationType("relProcessToApp", "BusinessProcess", "Application", "uses"),
        MetamodelRelationType(
            "relOrgToApp",
            "Organization",
            "Application",
            "uses",
            attributes_schema=[{"key": "usageType", "type": "single_select"}],
        ),
        MetamodelRelationType("relServerToApp", "Server", "Application", "hosts"),
    ]
    return MigrationSnapshot(
        version="turbo-ea",
        entities=entities,
        relations=relations,
        subscriptions=subscriptions,
        tags=tags,
        documents=documents,
        comments=comments,
        users=[],
        metamodel_types=metamodel_types,
        metamodel_relation_types=metamodel_relation_types,
    )


def _rows(wb, sheet: str) -> list[tuple]:
    return list(wb[sheet].iter_rows(values_only=True))


def _records(wb, sheet: str) -> list[dict]:
    rows = _rows(wb, sheet)
    keys = rows[0]
    return [dict(zip(keys, r)) for r in rows[2:]]


# ---------------------------------------------------------------------------
# Mapping tables — export and import must agree
# ---------------------------------------------------------------------------


class TestMappingConsistency:
    def test_every_export_type_maps_back_to_its_own_key(self):
        for tea_key, lx_name in EXPORT_TYPE_MAPPING.items():
            assert TYPE_MAPPING[lx_name] == tea_key, (tea_key, lx_name)

    def test_every_export_relation_maps_back_to_its_own_key(self):
        for tea_key, lx_name in EXPORT_RELATION_MAPPING.items():
            if lx_name == tea_key:
                continue  # relProcessDependency: exported under its Turbo EA key
            assert RELATION_MAPPING[lx_name] == tea_key, (tea_key, lx_name)

    def test_export_flip_set_is_the_mirror_of_the_import_flip_set(self):
        for tea_key in EXPORT_FLIP_DIRECTION:
            assert EXPORT_RELATION_MAPPING[tea_key] in FLIP_DIRECTION, tea_key
        flipped_lx = {EXPORT_RELATION_MAPPING[k] for k in EXPORT_FLIP_DIRECTION}
        for lx_name in FLIP_DIRECTION:
            if lx_name in RELATION_MAPPING:
                assert lx_name in flipped_lx or not lx_name.endswith("SuccessorRelation")

    def test_every_seeded_relation_type_has_an_export_name(self):
        from app.services.seed import RELATIONS

        seeded = {r["key"] for r in RELATIONS}
        missing = seeded - set(EXPORT_RELATION_MAPPING) - {"relProcessCalls"}
        assert not missing, missing


# ---------------------------------------------------------------------------
# Workbook shape
# ---------------------------------------------------------------------------


class TestWorkbookLayout:
    def test_payload_is_an_xlsx_and_the_sheet_set_mirrors_a_full_snapshot(self):
        data = build_workbook(_snapshot())
        assert is_xlsx_payload(data[:4])
        wb = load_workbook(BytesIO(data))
        names = wb.sheetnames
        assert names[0] == "ReadMe"
        assert names[-5:] == ["TagGroups", "Tags", "Documents", "Comments", "Types"]
        # Card types are named the LeanIX way; a custom type keeps its key.
        assert {"Application", "BusinessCapability", "Process", "UserGroup", "Server"} <= set(names)
        assert "BusinessProcess" not in names and "Organization" not in names
        assert "childParentRelation" in names
        assert "applicationSuccessorRelation" in names
        assert "processApplicationRelation" in names
        assert "relServerToApp" in names  # custom relation type keeps its key

    def test_every_sheet_has_a_key_row_and_a_label_row(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        for name in wb.sheetnames:
            if name == "ReadMe":
                continue
            rows = _rows(wb, name)
            assert len(rows) >= 2, name
            assert rows[0][0] in ("id", "factSheet", "factSheetType"), name
            assert rows[1][0] not in (None, rows[0][0]), name

    def test_fact_sheet_row_carries_core_lifecycle_tag_and_subscription_columns(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        apps = {r["id"]: r for r in _records(wb, "Application")}
        sf = apps["app-1"]
        assert sf["type"] == "Application"
        assert sf["name"] == "Salesforce"
        assert sf["displayName"] == "Salesforce"
        assert sf["status"] == "ACTIVE"
        assert sf["category"] == "businessApplication"
        assert sf["completion"] == 0.85
        assert sf["qualitySeal"] == "APPROVED"
        assert sf["lifecycle:plan"] == datetime(2019, 1, 1)
        assert sf["lifecycle:active"] == datetime(2020, 1, 1)
        assert sf["tags:Region"] == "EMEA"
        assert sf["tags:Stage"] == "Pilot"
        assert sf["subscriptions:RESPONSIBLE:Responsible"] == "owner@example.com"
        assert sf["subscriptions:OBSERVER:Observer"] == "a@example.com; b@example.com"
        assert sf["costTotalAnnual"] == 1200.5
        assert sf["hostingType"] == "cloud"
        assert sf["regions"] == "emea, apac"
        assert sf["externalId"] == "LX-42"
        # The empty card has the same columns, all blank.
        assert apps["app-2"]["tags:Region"] is None
        assert apps["app-2"]["lifecycle:plan"] is None

    def test_label_row_uses_field_labels(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        rows = _rows(wb, "Application")
        labels = dict(zip(rows[0], rows[1]))
        assert labels["costTotalAnnual"] == "Total annual cost"
        assert labels["lifecycle:plan"] == "Lifecycle: plan"
        assert labels["tags:Region"] == "Tags: Region"

    def test_hierarchy_becomes_child_parent_rows_with_path_display_names(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        rows = _records(wb, "childParentRelation")
        assert len(rows) == 1
        row = rows[0]
        assert row["type"] == "childParentRelation"
        assert row["fromRelatedFactSheetDisplayName"] == "Sales / Lead Mgmt"
        assert row["fromRelatedFactSheetType"] == "BusinessCapability"
        assert row["toRelatedFactSheetDisplayName"] == "Sales"
        assert row["toRelatedFactSheetType"] == "BusinessCapability"

    def test_successor_endpoints_are_swapped_to_leanix_direction(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        row = _records(wb, "applicationSuccessorRelation")[0]
        # LeanIX: from = the older card ("Salesforce has successor New CRM").
        assert row["fromRelatedFactSheetDisplayName"] == "Salesforce"
        assert row["toRelatedFactSheetDisplayName"] == "New CRM"

    def test_relation_row_translates_endpoint_types_and_carries_attributes(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        row = _records(wb, "applicationUserGroupRelation")[0]
        assert row["fromRelatedFactSheetType"] == "UserGroup"
        assert row["toRelatedFactSheetType"] == "Application"
        assert row["usageType"] == "owner"
        assert row["description"] == "Owning unit"
        proc = _records(wb, "processApplicationRelation")[0]
        assert proc["fromRelatedFactSheetType"] == "Process"

    def test_tag_sheets_follow_the_leanix_quirks(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        groups = {r["name"]: r for r in _records(wb, "TagGroups")}
        assert groups["Region"]["mode"] == "MULTIPLE"
        assert groups["Stage"]["mode"] == "SINGLE"
        assert groups["Stage"]["restrictToFactSheetTypes"] == "Application"
        tags = {r["id"]: r for r in _records(wb, "Tags")}
        # ``tagGroupId`` holds the group NAME in a LeanIX export.
        assert tags["tag-emea"]["tagGroupId"] == "Region"
        assert tags["tag-emea"]["backgroundColor"] == "#ff0000"

    def test_documents_and_comments_reference_cards_by_display_name_and_type(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        doc = _records(wb, "Documents")[0]
        assert doc["factSheet"] == "Salesforce"
        assert doc["type"] == "Application"
        assert doc["url"] == "https://wiki.example.com/sf"
        cmt = _records(wb, "Comments")[0]
        assert cmt["factSheet"] == "Salesforce"
        assert cmt["message"] == "Renewal due in Q3"
        assert cmt["userEmail"] == "owner@example.com"
        assert cmt["createdAt"] == datetime(2024, 3, 1, 12, 0)

    def test_types_and_readme_sheets_document_select_fields(self):
        wb = load_workbook(BytesIO(build_workbook(_snapshot())))
        types = {(r[0], r[1]): [v for v in r[2:] if v] for r in _rows(wb, "Types")[2:]}
        assert types[("Application", "hostingType")] == ["cloud", "onprem"]
        assert types[("Application", "regions")] == ["emea", "apac"]
        assert types[("Application", "category")] == ["businessApplication", "microservice"]
        assert types[("UserGroup", "category")] == ["team"]

        readme = _rows(wb, "ReadMe")
        by_col = {}
        scope = None
        for row in readme:
            if row[0] and not any(row[1:]):
                scope = row[0]
            elif row[0] and row[1]:
                by_col[(scope, row[0])] = row
        app_scope = "Fact Sheet type Application"
        assert by_col[(app_scope, "costTotalAnnual")][1] == "Money"
        assert by_col[(app_scope, "hostingType")][1] == "String"
        assert "Possible values: one of cloud, onprem." in by_col[(app_scope, "hostingType")][4]
        assert by_col[(app_scope, "regions")][1] == "String list"
        assert by_col[("All Fact Sheet types", "externalId")][1] == "String"

    def test_long_and_illegal_sheet_names_are_made_valid_and_unique(self):
        snap = _snapshot()
        snap.entities.append(SourceEntity("x-1", "A/Very:Long*Type?Name[That]Overflows", "x"))
        snap.entities.append(SourceEntity("x-2", "A/Very:Long*Type?Name[That]Overflowz", "y"))
        wb = load_workbook(BytesIO(build_workbook(snap)))
        long_titles = [n for n in wb.sheetnames if n.startswith("A_Very_Long")]
        assert len(long_titles) == 2
        assert all(len(n) <= 31 for n in long_titles)
        assert len(set(long_titles)) == 2
        # The canonical name survives inside the sheet.
        for title in long_titles:
            assert _records(wb, title)[0]["type"].startswith("A/Very:Long")

    def test_empty_workspace_still_produces_a_readable_workbook(self):
        snap = MigrationSnapshot("turbo-ea", [], [], [], [], [], [], [], [], [])
        data = build_workbook(snap)
        wb = load_workbook(BytesIO(data))
        assert wb.sheetnames == ["ReadMe", "TagGroups", "Tags", "Documents", "Comments", "Types"]
        parsed = parse_xlsx(BytesIO(data))
        assert parsed.entities == [] and parsed.parse_errors == []


# ---------------------------------------------------------------------------
# Round trip through the importer's parser + adapter tables
# ---------------------------------------------------------------------------


class TestRoundTrip:
    def test_export_then_parse_yields_the_same_landscape(self):
        snap = _snapshot()
        source = LeanixSource()
        parsed = parse_xlsx(BytesIO(build_workbook(snap)))
        assert parsed.parse_errors == []

        # Cards: same ids, Turbo EA types after the adapter's type map, same parents.
        by_id = {e.source_id: e for e in parsed.entities}
        assert set(by_id) == {e.source_id for e in snap.entities}
        for original in snap.entities:
            back = by_id[original.source_id]
            assert source.type_mapping.get(back.type, back.type) == original.type
            assert back.name == original.name
            assert back.parent_id == original.parent_id
            assert back.category == original.category
            assert back.lifecycle == original.lifecycle
            assert set(back.tags) == set(original.tags)
            assert back.status == original.status
        sf = by_id["app-1"]
        assert sf.description == "CRM"
        assert sf.quality_seal == "APPROVED"
        assert sf.completion == 0.85
        assert sf.custom_fields["costTotalAnnual"] == 1200.5
        assert sf.custom_fields["hostingType"] == "cloud"
        assert sf.custom_fields["externalId"] == "LX-42"

        # Relations: same Turbo EA keys and, after the import flip, the same direction.
        def to_tea(rel):
            key = source.relation_mapping.get(rel.type, rel.type)
            src, tgt = rel.from_entity_id, rel.to_entity_id
            if rel.type in source.flip_direction:
                src, tgt = tgt, src
            return (key, src, tgt)

        assert sorted(to_tea(r) for r in parsed.relations) == sorted(
            (r.type, r.from_entity_id, r.to_entity_id) for r in snap.relations
        )
        org_rel = next(r for r in parsed.relations if r.type == "applicationUserGroupRelation")
        assert org_rel.attributes["usageType"] == "owner"

        # Subscriptions route back to the same stakeholder role keys.
        roles = sorted(
            (s.user_email, source.map_subscription_role(s.role_name, s.role_type))
            for s in parsed.subscriptions
        )
        assert roles == [
            ("a@example.com", "observer"),
            ("b@example.com", "observer"),
            ("owner@example.com", "responsible"),
        ]

        # Tags, documents and comments.
        assert {(t.source_id, t.name, t.group_name) for t in parsed.tags} == {
            (t.source_id, t.name, t.group_name) for t in snap.tags
        }
        assert [(d.entity_id, d.name, d.url) for d in parsed.documents] == [
            ("app-1", "Runbook", "https://wiki.example.com/sf")
        ]
        assert [(c.entity_id, c.author_email, c.body) for c in parsed.comments] == [
            ("app-1", "owner@example.com", "Renewal due in Q3")
        ]

    def test_select_fields_come_back_as_selects_with_the_full_option_list(self):
        parsed = parse_xlsx(BytesIO(build_workbook(_snapshot())))
        app = next(m for m in parsed.metamodel_types if m.name == "Application")
        fields = {f.key: f for f in app.fields}
        assert fields["hostingType"].data_type == "SINGLE_SELECT"
        assert [o["key"] for o in fields["hostingType"].options] == ["cloud", "onprem"]
        assert fields["regions"].data_type == "MULTIPLE_SELECT"
        assert fields["costTotalAnnual"].data_type == "MONEY"

    def test_custom_type_and_relation_surface_as_metamodel_rows(self):
        parsed = parse_xlsx(BytesIO(build_workbook(_snapshot())))
        assert any(m.name == "Server" for m in parsed.metamodel_types)
        custom = next(r for r in parsed.metamodel_relation_types if r.name == "relServerToApp")
        assert (custom.source_type, custom.target_type) == ("Server", "Application")
