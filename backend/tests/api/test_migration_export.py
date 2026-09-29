"""``GET /migration/export`` — the workspace in a source platform's own format.

Gated on ``admin.export_workspace`` (the same authority as the workspace
bundle export, since the file is the whole landscape), registry-driven
through ``source_key``, and declared ahead of ``GET /migration/{id}`` so
the literal path is not swallowed by the UUID route.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.services.migration.sources.leanix.xlsx_parser import is_xlsx_payload
from tests.conftest import (
    auth_headers,
    create_card,
    create_card_type,
    create_relation,
    create_relation_type,
    create_role,
    create_user,
)


@pytest.fixture
def sources_registry_snapshot():
    """Restore the adapter registry after a test registers a throwaway source."""
    from app.services.migration.registry import SOURCES

    before = dict(SOURCES)
    yield
    SOURCES.clear()
    SOURCES.update(before)


@pytest.fixture
async def export_env(db):
    await create_role(db, key="admin", label="Admin", permissions={"*": True})
    await create_role(
        db, key="migrator", label="Migrator", permissions={"admin.migrate": True}, is_system=False
    )
    await create_role(
        db,
        key="exporter",
        label="Exporter",
        permissions={"admin.export_workspace": True},
        is_system=False,
    )
    admin = await create_user(db, email="admin@test.com", role="admin")
    migrator = await create_user(db, email="migrator@test.com", role="migrator")
    exporter = await create_user(db, email="exporter@test.com", role="exporter")

    await create_card_type(db, key="Application", label="Application", built_in=True)
    await create_card_type(db, key="BusinessProcess", label="Business Process", built_in=True)
    await create_relation_type(
        db,
        key="relProcessToApp",
        label="uses",
        source_type_key="BusinessProcess",
        target_type_key="Application",
    )
    app = await create_card(db, card_type="Application", name="Salesforce")
    proc = await create_card(db, card_type="BusinessProcess", name="Order to Cash")
    archived = await create_card(db, card_type="Application", name="Old CRM", status="ARCHIVED")
    await create_relation(db, type_key="relProcessToApp", source_id=proc.id, target_id=app.id)
    return {"admin": admin, "migrator": migrator, "exporter": exporter, "archived": archived}


class TestSourcesListing:
    async def test_leanix_reports_export_support(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/sources", headers=auth_headers(export_env["admin"])
        )
        assert resp.status_code == 200
        leanix = next(s for s in resp.json() if s["key"] == "leanix")
        assert leanix["supports_export"] is True


class TestExportRoute:
    async def test_requires_the_workspace_export_permission(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "leanix"},
            headers=auth_headers(export_env["migrator"]),
        )
        assert resp.status_code == 403

    async def test_the_export_permission_alone_is_enough(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "leanix"},
            headers=auth_headers(export_env["exporter"]),
        )
        assert resp.status_code == 200

    async def test_returns_a_leanix_workbook_as_an_attachment(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "leanix"},
            headers=auth_headers(export_env["admin"]),
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        disposition = resp.headers["content-disposition"]
        assert 'filename="leanix_export_' in disposition and disposition.endswith('.xlsx"')
        assert is_xlsx_payload(resp.content[:4])

        wb = load_workbook(BytesIO(resp.content))
        assert "Application" in wb.sheetnames
        assert "Process" in wb.sheetnames  # BusinessProcess → LeanIX "Process"
        assert "processApplicationRelation" in wb.sheetnames
        names = [r[2] for r in wb["Application"].iter_rows(min_row=3, values_only=True)]
        assert names == ["Salesforce"]  # archived card left out by default

    async def test_include_archived_flag(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "leanix", "include_archived": "true"},
            headers=auth_headers(export_env["admin"]),
        )
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        rows = list(wb["Application"].iter_rows(min_row=3, values_only=True))
        assert sorted((r[2], r[4]) for r in rows) == [
            ("Old CRM", "ARCHIVED"),
            ("Salesforce", "ACTIVE"),
        ]

    async def test_unknown_source_is_404(self, client, db, export_env):
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "not-a-source"},
            headers=auth_headers(export_env["admin"]),
        )
        assert resp.status_code == 404

    async def test_source_without_export_is_400(
        self, client, db, export_env, sources_registry_snapshot
    ):
        from app.services.migration.registry import register_source

        class _ReadOnlySource:
            key = "readonly"
            label = "Read only"
            accepted_extensions = (".bin",)
            type_mapping: dict[str, str] = {}
            relation_mapping: dict[str, str] = {}
            flip_direction: frozenset[str] = frozenset()
            field_type_mapping: dict[str, str] = {}
            auto_mapped_columns: tuple[tuple[str, str], ...] = ()

            def validate_payload(self, head: bytes) -> bool:
                return True

            def parse(self, path):
                raise NotImplementedError

            def post_build_card_payload(self, entity, target_type, payload):
                return None

            def map_subscription_role(self, role_name, role_type):
                return "responsible"

        register_source(_ReadOnlySource())
        resp = await client.get(
            "/api/v1/migration/export",
            params={"source_key": "readonly"},
            headers=auth_headers(export_env["admin"]),
        )
        assert resp.status_code == 400
        listed = await client.get(
            "/api/v1/migration/sources", headers=auth_headers(export_env["admin"])
        )
        assert next(s for s in listed.json() if s["key"] == "readonly")["supports_export"] is False
