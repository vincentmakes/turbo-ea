"""LeanIX xlsx (workspace export) writer — the mirror of ``xlsx_parser.py``.

Turns a :class:`MigrationSnapshot` built from the Turbo EA inventory
(:func:`app.services.migration.export_snapshot.build_export_snapshot`,
Turbo EA vocabulary) into the multi-sheet "Full Snapshot" workbook
LeanIX exports and the importer's parser reads back. The two modules
agree on every convention, and ``test_migration_leanix_xlsx_exporter.py``
round-trips a workbook through :func:`parse_xlsx` to keep it that way:

- Row 1 of every sheet holds the LeanIX field keys, row 2 the human
  labels, data starts on row 3.
- One sheet per fact-sheet type (``Application``, ``Process``, …) with
  the core columns, ``lifecycle:<phase>``, ``tags:<Group>`` (comma
  separated) and ``subscriptions:<RoleType>:<RoleName>`` (semicolon
  separated) columns, then one column per custom attribute.
- One sheet per relation type (``childParentRelation`` for the
  hierarchy, ``applicationITComponentRelation``, …) whose endpoints are
  ``(displayName, factSheetType)`` pairs — ``displayName`` is the
  hierarchy path joined with `` / ``, which is unique per type.
- Auxiliary sheets ``TagGroups``, ``Tags`` (``tagGroupId`` holds the
  group *name*), ``Documents``, ``Comments``, ``Types`` (observed enum
  values) and ``ReadMe`` (the field reference the importer trusts for
  field types and complete option lists).

Names are translated through the export tables in ``mappings.py``;
anything without an entry — an admin-created card type or relation
type — keeps its Turbo EA key, which the importer then surfaces as a
metamodel row for the admin to map.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]

from app.services.migration.snapshot import (
    MetamodelField,
    MigrationSnapshot,
    Relation,
    SourceEntity,
)
from app.services.migration.sources.leanix.mappings import (
    EXPORT_FLIP_DIRECTION,
    EXPORT_RELATION_MAPPING,
    EXPORT_TYPE_MAPPING,
)

HIERARCHY_SHEET = "childParentRelation"

# Core fact-sheet columns, in LeanIX's order, with their row-2 labels.
_FS_CORE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("id", "ID"),
    ("type", "Type"),
    ("name", "Name"),
    ("displayName", "Display Name"),
    ("status", "Status"),
    ("description", "Description"),
    ("category", "Category"),
    ("completion", "Completion"),
    ("qualitySeal", "Quality Seal"),
    ("createdAt", "Created at"),
    ("updatedAt", "Updated at"),
)

_REL_CORE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("id", "ID"),
    ("type", "Type"),
    ("fromRelatedFactSheetDisplayName", "From: Display Name"),
    ("fromRelatedFactSheetType", "From: Type"),
    ("toRelatedFactSheetDisplayName", "To: Display Name"),
    ("toRelatedFactSheetType", "To: Type"),
    ("status", "Status"),
    ("activeFrom", "Active from"),
    ("activeUntil", "Active until"),
)

# Lifecycle phases in chronological order — the column order LeanIX uses.
_LIFECYCLE_PHASES = ("plan", "phaseIn", "active", "phaseOut", "endOfLife")

# Turbo EA ``fields_schema`` type → the ``Type`` word the ReadMe sheet
# uses (the parser's ``_README_TYPE_MAP`` reads these back).
_README_TYPES: dict[str, str] = {
    "text": "String",
    "multiline_text": "String",
    "number": "Double",
    "cost": "Money",
    "percentage": "Percent",
    "boolean": "Boolean",
    "date": "Date",
    "url": "URL",
    "single_select": "String",
    "multiple_select": "String list",
}

# Columns every fact-sheet type shares; documented once under the
# ReadMe's "All Fact Sheet types" section.
_README_GLOBAL_ROWS: tuple[tuple[str, str, str], ...] = (
    ("id", "String (UUID)", "Turbo EA card id."),
    ("type", "String", "Fact Sheet type."),
    ("name", "String", "Card name."),
    ("displayName", "String", "Hierarchy path, ancestors joined with ' / '."),
    ("status", "String", "Possible values: one of ACTIVE, ARCHIVED."),
    ("description", "String", ""),
    ("category", "String", "Subtype key."),
    ("completion", "Percent", "Data-quality score, 0..1."),
    (
        "qualitySeal",
        "String",
        "Possible values: one of APPROVED, BROKEN_QUALITY_SEAL, DRAFT.",
    ),
    ("createdAt", "Datetime", ""),
    ("updatedAt", "Datetime", ""),
    ("externalId", "String", "External id carried by the card, when set."),
    ("alias", "String", "Card alias, when set."),
    ("lifecycle:plan", "Date", ""),
    ("lifecycle:phaseIn", "Date", ""),
    ("lifecycle:active", "Date", ""),
    ("lifecycle:phaseOut", "Date", ""),
    ("lifecycle:endOfLife", "Date", ""),
)

_SHEET_TITLE_BAD = re.compile(r"[\[\]:*?/\\]")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def build_workbook(snapshot: MigrationSnapshot) -> bytes:
    """Serialise ``snapshot`` as a LeanIX Full Snapshot workbook."""
    plan = _ExportPlan(snapshot)
    wb = Workbook(write_only=True)

    # Claim every sheet title up front so the ReadMe's "Relations to
    # Sheets Mapping" tail names the titles actually used.
    titles = _SheetTitles()
    for fixed in ("ReadMe", "TagGroups", "Tags", "Documents", "Comments", "Types"):
        titles.claim(fixed)
    type_titles = {lx: titles.claim(lx) for lx in plan.type_order}
    hierarchy_rows = plan.hierarchy_rows()
    hierarchy_title = titles.claim(HIERARCHY_SHEET) if hierarchy_rows else None
    relation_titles = {lx: titles.claim(lx) for lx in plan.relation_order}

    _write_readme(wb.create_sheet("ReadMe"), plan, relation_titles)

    for lx_type in plan.type_order:
        _write_fact_sheet_sheet(wb.create_sheet(type_titles[lx_type]), lx_type, plan)

    if hierarchy_title:
        ws = wb.create_sheet(hierarchy_title)
        _write_relation_sheet(ws, hierarchy_rows, attribute_keys=[])

    for lx_rel_type in plan.relation_order:
        rows = plan.relation_rows[lx_rel_type]
        ws = wb.create_sheet(relation_titles[lx_rel_type])
        _write_relation_sheet(ws, rows, attribute_keys=plan.relation_attribute_keys[lx_rel_type])

    _write_tag_groups(wb.create_sheet("TagGroups"), plan)
    _write_tags(wb.create_sheet("Tags"), plan)
    _write_documents(wb.create_sheet("Documents"), plan)
    _write_comments(wb.create_sheet("Comments"), plan)
    _write_types(wb.create_sheet("Types"), plan)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Plan — everything the sheets need, computed once
# ---------------------------------------------------------------------------


class _ExportPlan:
    def __init__(self, snapshot: MigrationSnapshot) -> None:
        self.snapshot = snapshot
        self.entity_by_id: dict[str, SourceEntity] = {e.source_id: e for e in snapshot.entities}
        self.tag_by_id = {t.source_id: t for t in snapshot.tags}
        self.fields_by_type: dict[str, list[MetamodelField]] = {
            mt.name: list(mt.fields) for mt in snapshot.metamodel_types
        }
        self.subtypes_by_type: dict[str, list[str]] = {
            mt.name: list(mt.subtypes) for mt in snapshot.metamodel_types
        }
        self.relation_schema: dict[str, list[dict[str, Any]]] = {
            rt.name: list(rt.attributes_schema or []) for rt in snapshot.metamodel_relation_types
        }

        # Entities grouped by their LeanIX type, in metamodel order.
        self.entities_by_lx_type: dict[str, list[SourceEntity]] = defaultdict(list)
        self.tea_types_of_lx_type: dict[str, list[str]] = defaultdict(list)
        for entity in snapshot.entities:
            lx = lx_type_name(entity.type)
            self.entities_by_lx_type[lx].append(entity)
            if entity.type not in self.tea_types_of_lx_type[lx]:
                self.tea_types_of_lx_type[lx].append(entity.type)
        metamodel_order = [mt.name for mt in snapshot.metamodel_types]
        ordered: list[str] = []
        for tea_type in metamodel_order + sorted(
            {e.type for e in snapshot.entities} - set(metamodel_order)
        ):
            lx = lx_type_name(tea_type)
            if lx in self.entities_by_lx_type and lx not in ordered:
                ordered.append(lx)
        self.type_order = ordered

        self.subscriptions_by_entity: dict[str, list] = defaultdict(list)
        for sub in snapshot.subscriptions:
            self.subscriptions_by_entity[sub.entity_id].append(sub)

        # Relations grouped by LeanIX relation name, endpoints already
        # oriented the LeanIX way and attribute keys collected per type.
        self.relation_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
        self.relation_attribute_keys: dict[str, list[str]] = {}
        seen_attr_keys: dict[str, list[str]] = defaultdict(list)
        for rel in snapshot.relations:
            row = self._relation_row(rel)
            if row is None:
                continue
            lx_rel = row["type"]
            self.relation_rows[lx_rel].append(row)
            for key in row["attributes"]:
                if key not in seen_attr_keys[lx_rel]:
                    seen_attr_keys[lx_rel].append(key)
        rel_order = [rt.name for rt in snapshot.metamodel_relation_types]
        ordered_rels: list[str] = []
        for tea_key in rel_order + sorted({r.type for r in snapshot.relations} - set(rel_order)):
            lx = lx_relation_name(tea_key)
            if lx in self.relation_rows and lx not in ordered_rels:
                ordered_rels.append(lx)
        self.relation_order = ordered_rels
        for lx_rel in ordered_rels:
            tea_keys = [k for k in rel_order if lx_relation_name(k) == lx_rel]
            schema_keys = [
                f.get("key")
                for k in tea_keys
                for f in self.relation_schema.get(k, [])
                if isinstance(f, dict) and f.get("key")
            ]
            extra = sorted(set(seen_attr_keys[lx_rel]) - set(schema_keys))
            self.relation_attribute_keys[lx_rel] = [
                k for k in schema_keys if k in seen_attr_keys[lx_rel]
            ] + extra

    # ---- relations -----------------------------------------------------

    def _relation_row(self, rel: Relation) -> dict[str, Any] | None:
        src = self.entity_by_id.get(rel.from_entity_id)
        tgt = self.entity_by_id.get(rel.to_entity_id)
        if src is None or tgt is None:
            return None
        if rel.type in EXPORT_FLIP_DIRECTION:
            src, tgt = tgt, src
        attributes = {
            k: v
            for k, v in (rel.attributes or {}).items()
            if k not in ("activeFrom", "activeUntil")
        }
        return {
            "id": rel.source_id,
            "type": lx_relation_name(rel.type),
            "from_name": display_name(src),
            "from_type": lx_type_name(src.type),
            "to_name": display_name(tgt),
            "to_type": lx_type_name(tgt.type),
            "status": "ACTIVE",
            "activeFrom": (rel.attributes or {}).get("activeFrom"),
            "activeUntil": (rel.attributes or {}).get("activeUntil"),
            "attributes": attributes,
        }

    def hierarchy_rows(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for entity in self.snapshot.entities:
            if not entity.parent_id:
                continue
            parent = self.entity_by_id.get(entity.parent_id)
            if parent is None:
                continue
            rows.append(
                {
                    "id": f"{HIERARCHY_SHEET}:{entity.source_id}",
                    "type": HIERARCHY_SHEET,
                    "from_name": display_name(entity),
                    "from_type": lx_type_name(entity.type),
                    "to_name": display_name(parent),
                    "to_type": lx_type_name(parent.type),
                    "status": "ACTIVE",
                    "activeFrom": None,
                    "activeUntil": None,
                    "attributes": {},
                }
            )
        return rows

    # ---- per-type column planning ---------------------------------------

    def tag_groups_for(self, lx_type: str) -> list[str]:
        """Tag-group column order: groups open to the type, then any other group a card uses."""
        tea_types = set(self.tea_types_of_lx_type[lx_type])
        names: list[str] = []
        for tag in self.snapshot.tags:
            if not tag.group_name or tag.group_name in names:
                continue
            allowed = tag.group_restrict_to_types
            if not allowed or tea_types & set(allowed):
                names.append(tag.group_name)
        for entity in self.entities_by_lx_type[lx_type]:
            for tag_id in entity.tags:
                tag = self.tag_by_id.get(tag_id)
                if tag and tag.group_name and tag.group_name not in names:
                    names.append(tag.group_name)
        return names

    def subscription_columns_for(self, lx_type: str) -> list[tuple[str, str]]:
        cols: list[tuple[str, str]] = []
        for entity in self.entities_by_lx_type[lx_type]:
            for sub in self.subscriptions_by_entity.get(entity.source_id, []):
                key = (sub.role_type or "RESPONSIBLE", sub.role_name or "")
                if key not in cols:
                    cols.append(key)
        return sorted(cols)

    def attribute_columns_for(self, lx_type: str) -> list[tuple[str, str]]:
        """``(key, label)`` — schema order first, then any other key a card carries, sorted."""
        cols: list[tuple[str, str]] = []
        seen: set[str] = set()
        for tea_type in self.tea_types_of_lx_type[lx_type]:
            for f in self.fields_by_type.get(tea_type, []):
                if f.key not in seen:
                    seen.add(f.key)
                    cols.append((f.key, f.label or f.key))
        extra: set[str] = set()
        for entity in self.entities_by_lx_type[lx_type]:
            extra.update(k for k in entity.custom_fields if k not in seen)
        used = {k for e in self.entities_by_lx_type[lx_type] for k in e.custom_fields}
        cols = [c for c in cols if c[0] in used]
        cols.extend((k, k) for k in sorted(extra))
        return cols

    def lifecycle_phases_for(self, lx_type: str) -> list[str]:
        used = {p for e in self.entities_by_lx_type[lx_type] for p in e.lifecycle}
        known = [p for p in _LIFECYCLE_PHASES if p in used]
        return known + sorted(used - set(_LIFECYCLE_PHASES))

    def select_fields(self) -> list[tuple[str, MetamodelField]]:
        """``(lx_type, field)`` for every select field of an exported type, plus ``category``."""
        out: list[tuple[str, MetamodelField]] = []
        for lx_type in self.type_order:
            for tea_type in self.tea_types_of_lx_type[lx_type]:
                subtypes = self.subtypes_by_type.get(tea_type) or []
                if subtypes:
                    out.append(
                        (
                            lx_type,
                            MetamodelField(
                                type_name=tea_type,
                                key="category",
                                label="Category",
                                data_type="single_select",
                                options=[{"key": s, "label": s} for s in subtypes],
                            ),
                        )
                    )
                for f in self.fields_by_type.get(tea_type, []):
                    if f.data_type in ("single_select", "multiple_select") and f.options:
                        out.append((lx_type, f))
        return out


# ---------------------------------------------------------------------------
# Sheet writers
# ---------------------------------------------------------------------------


def _write_fact_sheet_sheet(ws: Any, lx_type: str, plan: _ExportPlan) -> None:
    phases = plan.lifecycle_phases_for(lx_type)
    tag_groups = plan.tag_groups_for(lx_type)
    sub_cols = plan.subscription_columns_for(lx_type)
    attr_cols = plan.attribute_columns_for(lx_type)

    keys = [k for k, _ in _FS_CORE_COLUMNS]
    labels = [label for _, label in _FS_CORE_COLUMNS]
    for phase in phases:
        keys.append(f"lifecycle:{phase}")
        labels.append(f"Lifecycle: {phase}")
    for group in tag_groups:
        keys.append(f"tags:{group}")
        labels.append(f"Tags: {group}")
    for role_type, role_name in sub_cols:
        keys.append(f"subscriptions:{role_type}:{role_name}")
        labels.append(f"Subscriptions: {role_type}: {role_name}")
    for key, label in attr_cols:
        keys.append(key)
        labels.append(label)
    ws.append(keys)
    ws.append(labels)

    for entity in plan.entities_by_lx_type[lx_type]:
        row: list[Any] = [
            entity.source_id,
            lx_type,
            entity.name,
            display_name(entity),
            entity.status or "ACTIVE",
            entity.description,
            entity.category,
            entity.completion,
            entity.quality_seal,
            entity.raw.get("createdAt"),
            entity.raw.get("updatedAt"),
        ]
        for phase in phases:
            row.append(_cell_date(entity.lifecycle.get(phase)))
        tags_by_group: dict[str, list[str]] = defaultdict(list)
        for tag_id in entity.tags:
            tag = plan.tag_by_id.get(tag_id)
            if tag and tag.group_name:
                tags_by_group[tag.group_name].append(tag.name)
        for group in tag_groups:
            row.append(", ".join(tags_by_group.get(group, [])) or None)
        emails_by_col: dict[tuple[str, str], list[str]] = defaultdict(list)
        for sub in plan.subscriptions_by_entity.get(entity.source_id, []):
            if sub.user_email:
                emails_by_col[(sub.role_type or "RESPONSIBLE", sub.role_name or "")].append(
                    sub.user_email
                )
        for col in sub_cols:
            row.append("; ".join(emails_by_col.get(col, [])) or None)
        for key, _ in attr_cols:
            row.append(_cell_value(entity.custom_fields.get(key)))
        ws.append(row)


def _write_relation_sheet(
    ws: Any, rows: list[dict[str, Any]], *, attribute_keys: list[str]
) -> None:
    keys = [k for k, _ in _REL_CORE_COLUMNS] + list(attribute_keys)
    labels = [label for _, label in _REL_CORE_COLUMNS] + list(attribute_keys)
    ws.append(keys)
    ws.append(labels)
    for row in rows:
        ws.append(
            [
                row["id"],
                row["type"],
                row["from_name"],
                row["from_type"],
                row["to_name"],
                row["to_type"],
                row["status"],
                _cell_value(row["activeFrom"]),
                _cell_value(row["activeUntil"]),
            ]
            + [_cell_value(row["attributes"].get(k)) for k in attribute_keys]
        )


def _write_tag_groups(ws: Any, plan: _ExportPlan) -> None:
    ws.append(["id", "name", "mode", "restrictToFactSheetTypes"])
    ws.append(["ID", "Name", "Mode", "Restrict to Fact Sheet types"])
    seen: set[str] = set()
    for tag in plan.snapshot.tags:
        if not tag.group_name or tag.group_name in seen:
            continue
        seen.add(tag.group_name)
        restrict = ", ".join(lx_type_name(t) for t in (tag.group_restrict_to_types or []))
        ws.append([tag.group_name, tag.group_name, tag.group_mode or "MULTIPLE", restrict or None])


def _write_tags(ws: Any, plan: _ExportPlan) -> None:
    ws.append(["id", "name", "tagGroupId", "backgroundColor"])
    ws.append(["ID", "Name", "Tag group", "Background color"])
    for tag in plan.snapshot.tags:
        if not tag.group_name:
            continue
        ws.append([tag.source_id, tag.name, tag.group_name, tag.color])


def _write_documents(ws: Any, plan: _ExportPlan) -> None:
    ws.append(["id", "factSheet", "type", "name", "url"])
    ws.append(["ID", "Fact Sheet", "Fact Sheet type", "Name", "URL"])
    for doc in plan.snapshot.documents:
        entity = plan.entity_by_id.get(doc.entity_id)
        if entity is None:
            continue
        ws.append(
            [doc.source_id, display_name(entity), lx_type_name(entity.type), doc.name, doc.url]
        )


def _write_comments(ws: Any, plan: _ExportPlan) -> None:
    ws.append(["factSheet", "type", "message", "userEmail", "createdAt"])
    ws.append(["Fact Sheet", "Fact Sheet type", "Message", "User", "Created at"])
    for comment in plan.snapshot.comments:
        entity = plan.entity_by_id.get(comment.entity_id)
        if entity is None or not comment.body:
            continue
        ws.append(
            [
                display_name(entity),
                lx_type_name(entity.type),
                comment.body,
                comment.author_email,
                _cell_naive(comment.created_at),
            ]
        )


def _write_types(ws: Any, plan: _ExportPlan) -> None:
    select_fields = plan.select_fields()
    width = max((len(f.options) for _, f in select_fields), default=0)
    ws.append(["factSheetType", "fieldName"] + [f"value{i + 1}" for i in range(width)])
    ws.append(["Fact Sheet type", "Field"] + [f"Value {i + 1}" for i in range(width)])
    for lx_type, f in select_fields:
        ws.append([lx_type, f.key] + [o.get("key") for o in f.options])


def _write_readme(ws: Any, plan: _ExportPlan, relation_titles: dict[str, str]) -> None:
    ws.append(["Turbo EA workspace export in the SAP LeanIX Full Snapshot layout."])
    ws.append(
        [
            "Row 1 of every sheet holds the field keys, row 2 the labels; "
            "data starts on row 3. Relation endpoints reference fact sheets by "
            "(displayName, type)."
        ]
    )
    ws.append([])
    ws.append(["Column header", "Type", "Mandatory", "Read only", "Remarks"])
    ws.append(["All Fact Sheet types"])
    for col, lx_type, remarks in _README_GLOBAL_ROWS:
        ws.append([col, lx_type, "mandatory" if col in ("id", "type", "name") else "", "", remarks])
    for lx_type in plan.type_order:
        ws.append([f"Fact Sheet type {lx_type}"])
        for tea_type in plan.tea_types_of_lx_type[lx_type]:
            subtypes = plan.subtypes_by_type.get(tea_type) or []
            if subtypes:
                ws.append(
                    [
                        "category",
                        "String",
                        "",
                        "",
                        "Possible values: one of " + ", ".join(subtypes) + ".",
                    ]
                )
            for f in plan.fields_by_type.get(tea_type, []):
                ws.append(
                    [
                        f.key,
                        _README_TYPES.get(f.data_type, "String"),
                        "",
                        "",
                        _readme_remarks(f),
                    ]
                )
    ws.append(["Relations to Sheets Mapping"])
    for lx_rel in plan.relation_order:
        ws.append([lx_rel, relation_titles[lx_rel]])


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def lx_type_name(tea_type: str) -> str:
    return EXPORT_TYPE_MAPPING.get(tea_type, tea_type)


def lx_relation_name(tea_key: str) -> str:
    return EXPORT_RELATION_MAPPING.get(tea_key, tea_key)


def display_name(entity: SourceEntity) -> str:
    return entity.display_name or entity.name


def _readme_remarks(f: MetamodelField) -> str:
    if f.data_type not in ("single_select", "multiple_select") or not f.options:
        return ""
    values = [str(o.get("key")) for o in f.options if o.get("key") is not None]
    return "Possible values: one of " + ", ".join(values) + "." if values else ""


def _cell_value(value: Any) -> Any:
    """Flatten a JSON attribute value into something a spreadsheet cell holds."""
    if value is None or value == "":
        return None
    if isinstance(value, bool | int | float | str):
        return value
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, list | tuple | set):
        return ", ".join(str(v) for v in value if v not in (None, "")) or None
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _cell_date(value: Any) -> Any:
    """A lifecycle value is an ISO date string; write it as a real date cell."""
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return str(value)


def _cell_naive(value: datetime | None) -> datetime | None:
    # openpyxl refuses timezone-aware datetimes.
    return value.replace(tzinfo=None) if value is not None else None


class _SheetTitles:
    """Excel sheet titles: at most 31 chars, no ``[]:*?/\\``, unique per workbook."""

    def __init__(self) -> None:
        self._used: set[str] = set()

    @staticmethod
    def sanitise(name: str) -> str:
        return _SHEET_TITLE_BAD.sub("_", name)[:31] or "Sheet"

    def claim(self, name: str) -> str:
        base = self.sanitise(name)
        title = base
        n = 2
        while title.lower() in self._used:
            suffix = f"~{n}"
            title = base[: 31 - len(suffix)] + suffix
            n += 1
        self._used.add(title.lower())
        return title
